"""Дозаполняет индекс: корни корпуса из таблицы источников — страницы с вложениями, файлы, почтовый архив — и принятое через входящую папку.

Возобновляемый: пройденные файлы пишутся в ingested.txt, при перезапуске пропускаются.
Текст извлекается из PDF, docx, xlsx, msg, eml и простых текстовых форматов.
Картинки, архивы и бинарники пропускаются. Файл, который не прочитан из-за отсутствия библиотеки формата, в ingested.txt не пишется:
причина называется сообщением lib.missing, а после установки библиотеки повторный запуск его проиндексирует.
"""
import email, json, os, re, sys, traceback, urllib.request
from email import policy
import doctor
import ingest
import messages
import perms
import settings
import sources

_SETTINGS = settings.startup()
SOURCES = sources.startup()
_HOME = _SETTINGS["home"]
DB = os.path.join(_HOME, "index", "lance")
DONE = os.path.join(_HOME, "index", "ingested.txt")
OLLAMA, MODEL, DIM = _SETTINGS["embed_url"], _SETTINGS["embed_model"], _SETTINGS["embed_dim"]      # служба, модель и размерность векторов
MAX_CHARS, BATCH = 4000, 48
MAX_FILE_CHARS = 400_000          # обрезаем гигантские выгрузки

# корни обхода берутся из таблицы источников (sources.json каталога архива): вложения страниц раньше самих страниц, принятое
# через входящую папку последним — при полной пересборке индекса оно не должно пропасть. Без таблицы обходятся одни входящие.
CORPUS = os.path.join(_HOME, "corpus")
ROOTS = SOURCES.walk(CORPUS)
# _карантин лежит рядом в corpus, но в ROOTS его нет — секреты и личное в индекс не идут
PLAIN = {".txt", ".csv", ".tsv", ".md", ".log", ".sql", ".json", ".yaml", ".yml", ".properties", ".conf", ".ini"}


# ── извлечение текста ───────────────────────────────────────────
def from_pdf(p):
    import fitz
    with fitz.open(p) as d:
        return "\n".join(pg.get_text() for pg in d)


def from_docx(p):
    import docx
    d = docx.Document(p)
    parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def from_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
            if len("\n".join(out)) > MAX_FILE_CHARS:
                break
    wb.close()
    return "\n".join(out)


def from_msg(p):
    import extract_msg
    m = extract_msg.Message(p)
    head = f"От: {m.sender}\nКому: {m.to}\nТема: {m.subject}\nДата: {m.date}"
    body = m.body or ""
    m.close()
    return head + "\n\n" + body


def from_eml(p):
    with open(p, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    head = f"От: {msg.get('From')}\nКому: {msg.get('To')}\nТема: {msg.get('Subject')}\nДата: {msg.get('Date')}"
    body = ""
    try:
        part = msg.get_body(preferencelist=("plain", "html"))
        if part:
            body = part.get_content()
            if part.get_content_type() == "text/html":
                body = re.sub(r"<[^>]+>", " ", body)
    except Exception:
        pass
    return head + "\n\n" + body


def from_pptx(p):
    from pptx import Presentation
    out = []
    for i, sl in enumerate(Presentation(p).slides, 1):
        out.append(f"## Слайд {i}")
        for sh in sl.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def from_markup(p):
    """bpmn, svg, xml, html: оставляем текстовые узлы и подписи, теги выбрасываем."""
    with open(p, encoding="utf-8", errors="replace") as f:
        raw = f.read(MAX_FILE_CHARS)
    names = re.findall(r'\b(?:name|id|label|documentation)="([^"]{2,120})"', raw)
    text = re.sub(r"<[^>]+>", " ", raw)
    text = re.sub(r"\s+", " ", text)
    return "\n".join(dict.fromkeys(names)) + "\n" + text


EXTRACT = {".pdf": from_pdf, ".docx": from_docx, ".xlsx": from_xlsx, ".xlsm": from_xlsx,
           ".msg": from_msg, ".eml": from_eml, ".pptx": from_pptx,
           ".bpmn": from_markup, ".svg": from_markup, ".xml": from_markup,
           ".html": from_markup, ".htm": from_markup, ".drawio": from_markup}


# библиотеки разбора форматов (имена при импорте, как в doctor.LIBRARIES): без них файл не читается, и это не «битый файл»
FORMAT_MODULES = ("docx", "openpyxl", "pptx", "extract_msg", "pymupdf")
OLD_NAMES = {"fitz": "pymupdf"}                  # from_pdf берёт PyMuPDF под прежним именем


def missing_library(error):
    """Сообщение lib.missing, если файл не прочитан из-за отсутствия библиотеки формата (docx, xlsx, pptx, msg, pdf); иначе None. Не хватает зависимости
    самой библиотеки или это любая другая ошибка — не «нет библиотеки»: файл тогда ведёт себя как раньше (нечитаемый)."""
    if not isinstance(error, ModuleNotFoundError):
        return None
    top = (error.name or "").split(".")[0]
    top = OLD_NAMES.get(top, top)
    for module, package, must, use in doctor.LIBRARIES:
        if module == top and module in FORMAT_MODULES:
            return messages.make("lib.missing", package=package, use=use, file=doctor.REQUIRED_FILE if must else doctor.OPTIONAL_FILE)
    return None


def read_text(path, ext):
    if ext in EXTRACT:
        return EXTRACT[ext](path)
    if ext in PLAIN:
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read(MAX_FILE_CHARS)
    return None


# ── нарезка и эмбеддинги ────────────────────────────────────────
def chunks(body):
    body = body[:MAX_FILE_CHARS]
    if len(body) <= MAX_CHARS:
        return [body] if body.strip() else []
    out, p = [], body
    while len(p) > MAX_CHARS:
        out.append(p[:MAX_CHARS])
        p = p[MAX_CHARS - 200:]
    out.append(p)
    return [c for c in out if c.strip()]


def embed_options(env=None):
    """Векторы по умолчанию считаются на процессоре: видеопамять целиком отдана большой модели, запаса под модель векторов на карте нет.
    Плановая пересборка при выгруженной большой модели — по явному решению владельца: настройка embed_gpu. Настройки читаются при каждом вызове
    (из переданного окружения или окружения процесса)."""
    return {} if settings.load(env=env)["embed_gpu"] else {"num_gpu": 0}


def embed(texts):
    body = {"model": MODEL, "input": texts}
    options = embed_options()
    if options:
        body["options"] = options
    req = urllib.request.Request(OLLAMA, json.dumps(body).encode(), {"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=600) as r:
        return json.load(r)["embeddings"]


# ── обход ───────────────────────────────────────────────────────
def work_items(done):
    for source, root, needle in ROOTS:
        for dirpath, _, files in os.walk(root):
            if "_карантин" in dirpath.split(os.sep):
                continue              # карантин не индексируется никогда, где бы он ни лежал
            if needle:
                # "!x" — обходить всё, кроме x; иначе только то, где есть x
                if needle.startswith("!"):
                    if needle[1:] in dirpath:
                        continue
                elif needle not in dirpath:
                    continue
            for f in files:
                if f.startswith("._"):
                    continue
                ext = os.path.splitext(f)[1].lower()
                if ext not in EXTRACT and ext not in PLAIN:
                    continue
                if source == "mail" and ext == ".json":
                    continue          # рядом лежит .eml с тем же содержимым
                p = os.path.join(dirpath, f)
                if p in done:
                    continue
                yield source, root, p, ext


def meta_for(source, root, path, text):
    """(база, раздел, название, дата) документа: правила — в таблице источников, одни для индексатора, приёмки, чистки повторов и починки дат."""
    return SOURCES.meta_for(source, root, path, text)


def flush(tbl, buf, fdone):
    try:
        vecs = embed([r["text"] for r in buf])
    except Exception:
        buf2, vecs = [], []
        for r in buf:
            try:
                vecs.append(embed([r["text"]])[0])
                buf2.append(r)
            except Exception:
                print(f"пропуск чанка: {r['path']} #{r['chunk']}", flush=True)
        buf = buf2
        if not buf:
            return 0
    for r, v in zip(buf, vecs):
        r["vector"] = v
    tbl.add([{k: v for k, v in r.items() if k != "_src"} for r in buf])
    for p in dict.fromkeys(r["_src"] for r in buf):
        fdone.write(p + "\n")
    fdone.flush()
    return len(buf)


def main():
    lancedb = ingest.require_lancedb()           # библиотека индекса загружается здесь: без неё модуль всё равно загрузится, а команда ответит отказом
    done = set()
    if os.path.exists(DONE):
        with open(DONE, encoding="utf-8") as f:
            done = {l.rstrip("\n") for l in f}
    print(f"уже обработано файлов: {len(done)}", flush=True)

    tbl = lancedb.connect(DB).open_table("docs")
    n, files, buf = 0, 0, []
    lacking = {}                       # пакет -> сколько файлов не прочитано без него в этом запуске
    with open(DONE, "a", encoding="utf-8") as fdone:
        for source, root, path, ext in work_items(done):
            try:
                text = read_text(path, ext)
            except Exception as e:
                note = missing_library(e)
                if note is not None:
                    # в список пройденного такой файл не пишется: после установки библиотеки он должен попасть в индекс
                    if note.args["package"] not in lacking:
                        print(note, flush=True)
                    lacking[note.args["package"]] = lacking.get(note.args["package"], 0) + 1
                    continue
                print(f"нечитаемо ({type(e).__name__}): {path}", flush=True)
                fdone.write(path + "\n")
                continue
            files += 1
            if not text or not text.strip():
                fdone.write(path + "\n")
                continue
            src, space, title, updated = meta_for(source, root, path, text)
            rel = os.path.relpath(path, os.path.dirname(root))
            for i, c in enumerate(chunks(text)):
                buf.append({"path": rel, "source": src, "space": space, "title": title,
                            "updated": updated, "url": "", "chunk": i, "text": c, "_src": path})
                if len(buf) >= BATCH:
                    n += flush(tbl, buf, fdone)
                    buf = []
                    if n % 4800 == 0:
                        print(f"+{n} чанков, файлов {files}, сейчас {source}", flush=True)
        if buf:
            n += flush(tbl, buf, fdone)
    skipped = f", пропущено без библиотек: {sum(lacking.values())} (в список пройденного не записаны)" if lacking else ""
    print(f"наполнение закончено: +{n} чанков из {files} файлов{skipped}", flush=True)


def selftest():
    assert chunks("  ") == []
    assert chunks("коротко") == ["коротко"]
    big = "я" * 9000
    cs = chunks(big)
    assert len(cs) == 3 and all(len(c) <= MAX_CHARS for c in cs), [len(c) for c in cs]
    # правила баз — в таблице источников (их проверяют tests/test_sources*.py на выдуманных образцах); здесь — встроенный корень входящих
    root = os.path.join(os.sep, "tmp", "r", sources.INTAKE)
    sr, sp, ti, up = meta_for(sources.INTAKE, root, os.path.join(root, "b", "_attachments", "IT-123", "2023-08-22_file.pdf"), "")
    assert (sr, sp, ti, up) == (sources.INTAKE, "b", "IT-123 · 2023-08-22_file.pdf", "2023-08-22"), (sr, sp, ti, up)
    print("selftest ok")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
    else:
        perms.close_umask()                              # таблица индекса и список пройденного — только владельцу, при любой маске запустившего
        try:
            main()
        except messages.CodedError as e:                 # отказ с кодом (нет библиотеки индекса): одна строка, как у поиска
            print(f"ошибка: {messages.of(e)}", file=sys.stderr)
            sys.exit(1)
        except Exception:
            traceback.print_exc()
            sys.exit(1)
